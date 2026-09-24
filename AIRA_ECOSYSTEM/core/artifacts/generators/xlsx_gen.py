"""core/artifacts/generators/xlsx_gen.py - SpreadsheetSpec -> .xlsx (openpyxl)."""
from __future__ import annotations

from pathlib import Path

from core.artifacts.generators.base import GenerationError, ensure_parent
from core.artifacts.specs import SpreadsheetSpec

_NUMBER_FORMATS = {
    "number": "0.00",
    "integer": "0",
    "currency": '"Rp"#,##0.00',
    "percent": "0.00%",
    "date": "yyyy-mm-dd",
    "text": "@",
}


def generate_xlsx(spec: SpreadsheetSpec, output_path: Path) -> dict:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:
        raise GenerationError("openpyxl tidak terpasang (pip install openpyxl).") from exc

    ensure_parent(output_path)

    workbook = Workbook()
    workbook.remove(workbook.active)

    counts = {"worksheets": 0, "rows": 0}

    for sheet_spec in spec.worksheets:
        sheet = workbook.create_sheet(title=(sheet_spec.name or "Sheet")[:31])

        for col_index, header in enumerate(sheet_spec.headers, start=1):
            cell = sheet.cell(row=1, column=col_index, value=header)
            cell.font = Font(bold=True)

        column_index = {header: i + 1 for i, header in enumerate(sheet_spec.headers)}

        for row_offset, row in enumerate(sheet_spec.rows, start=2):
            for col_index, value in enumerate(row, start=1):
                sheet.cell(row=row_offset, column=col_index, value=_coerce_value(value))

            for header, fmt in sheet_spec.column_formats.items():
                idx = column_index.get(header)
                if idx and fmt in _NUMBER_FORMATS:
                    sheet.cell(row=row_offset, column=idx).number_format = _NUMBER_FORMATS[fmt]

            counts["rows"] += 1

        counts["worksheets"] += 1

    if not workbook.sheetnames:
        workbook.create_sheet(title="Sheet1")

    try:
        workbook.save(str(output_path))
    except OSError as exc:
        raise GenerationError(f"Gagal menyimpan .xlsx: {exc}") from exc

    return counts


def _coerce_value(value):
    if isinstance(value, str):
        stripped = value.strip()
        for caster in (int, float):
            try:
                return caster(stripped)
            except ValueError:
                continue
    return value