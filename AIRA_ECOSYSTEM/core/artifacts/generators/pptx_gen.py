"""core/artifacts/generators/pptx_gen.py - PresentationSpec -> .pptx (python-pptx)."""
from __future__ import annotations

from pathlib import Path

from core.artifacts.generators.base import GenerationError, ensure_parent
from core.artifacts.specs import PresentationSpec


def generate_pptx(spec: PresentationSpec, output_path: Path) -> dict:
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:
        raise GenerationError("python-pptx tidak terpasang (pip install python-pptx).") from exc

    ensure_parent(output_path)

    presentation = Presentation()
    title_layout = presentation.slide_layouts[0]
    content_layout = presentation.slide_layouts[1]

    title_slide = presentation.slides.add_slide(title_layout)
    title_slide.shapes.title.text = spec.title

    counts = {"sections": 0, "tables": 0, "image_placeholders": 0}

    for section in spec.sections:
        slide = presentation.slides.add_slide(content_layout)
        slide.shapes.title.text = section.heading

        body = slide.placeholders[1].text_frame
        body.clear()

        if section.content:
            body.text = section.content[0]
            for bullet in section.content[1:]:
                paragraph = body.add_paragraph()
                paragraph.text = bullet

        if section.image_ref:
            note_box = slide.shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(9), Inches(1))
            note_box.text_frame.text = f"[Gambar: {section.image_ref}]"
            counts["image_placeholders"] += 1

        if section.table:
            headers = section.table.get("headers", [])
            rows = section.table.get("rows", [])
            if headers:
                table_shape = slide.shapes.add_table(
                    len(rows) + 1, len(headers), Inches(0.5), Inches(4.2), Inches(9), Inches(1.2),
                ).table
                for col_index, header in enumerate(headers):
                    table_shape.cell(0, col_index).text = str(header)
                for row_index, row in enumerate(rows, start=1):
                    for col_index, value in enumerate(row):
                        table_shape.cell(row_index, col_index).text = str(value)
                counts["tables"] += 1

        counts["sections"] += 1

    if spec.conclusion:
        conclusion_slide = presentation.slides.add_slide(content_layout)
        conclusion_slide.shapes.title.text = "Kesimpulan"
        conclusion_slide.placeholders[1].text_frame.text = spec.conclusion

    try:
        presentation.save(str(output_path))
    except OSError as exc:
        raise GenerationError(f"Gagal menyimpan .pptx: {exc}") from exc

    return counts