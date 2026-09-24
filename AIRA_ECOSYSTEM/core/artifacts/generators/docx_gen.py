"""core/artifacts/generators/docx_gen.py - DocumentSpec -> .docx (python-docx)."""
from __future__ import annotations

from pathlib import Path

from core.artifacts.generators.base import GenerationError, ensure_parent
from core.artifacts.specs import DocumentSpec


def generate_docx(spec: DocumentSpec, output_path: Path) -> dict:
    try:
        from docx import Document
    except ImportError as exc:
        raise GenerationError("python-docx tidak terpasang (pip install python-docx).") from exc

    ensure_parent(output_path)

    document = Document()
    document.add_heading(spec.title, level=0)

    counts = {"headings": 0, "paragraphs": 0, "lists": 0, "tables": 0}

    for block in spec.blocks:
        if block.type == "heading":
            document.add_heading(block.text, level=max(1, min(block.level, 6)))
            counts["headings"] += 1

        elif block.type == "paragraph":
            para = document.add_paragraph()
            run = para.add_run(block.text)
            if block.style == "bold":
                run.bold = True
            elif block.style == "italic":
                run.italic = True
            counts["paragraphs"] += 1

        elif block.type == "list":
            style = "List Number" if block.ordered else "List Bullet"
            for item in block.items:
                document.add_paragraph(item, style=style)
            counts["lists"] += 1

        elif block.type == "table":
            table = document.add_table(rows=1, cols=len(block.headers))
            try:
                table.style = "Light Grid Accent 1"
            except KeyError:
                pass  # style not present in this template - table still works
            header_cells = table.rows[0].cells
            for i, header in enumerate(block.headers):
                header_cells[i].text = str(header)
            for row in block.rows:
                cells = table.add_row().cells
                for i, value in enumerate(row):
                    cells[i].text = str(value)
            counts["tables"] += 1

    try:
        document.save(str(output_path))
    except OSError as exc:
        raise GenerationError(f"Gagal menyimpan .docx: {exc}") from exc

    return counts