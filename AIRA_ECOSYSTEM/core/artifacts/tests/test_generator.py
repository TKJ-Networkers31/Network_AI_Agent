import csv
import tempfile
import unittest
from pathlib import Path

from core.artifacts.generators import (
    generate_csv, generate_docx, generate_markdown, generate_pdf,
    generate_pptx, generate_txt, generate_xlsx,
)
from core.artifacts.specs import (
    CsvSpec, DocumentSpec, PresentationSpec, SlideSection, SpreadsheetSpec, WorksheetSpec,
    heading, paragraph, bullet_list, table_block,
)


def _doc_spec() -> DocumentSpec:
    return DocumentSpec(title="Laporan Jaringan", blocks=[
        heading("Ringkasan", 1),
        paragraph("Semua router online."),
        bullet_list(["R1 online", "R2 online"]),
        table_block(["Device", "Status"], [["R1", "UP"], ["R2", "UP"]]),
    ])


def _sheet_spec() -> SpreadsheetSpec:
    return SpreadsheetSpec(worksheets=[
        WorksheetSpec(name="Traffic", headers=["Device", "Mbps"], rows=[["R1", "12.5"], ["R2", "8.2"]],
                      column_formats={"Mbps": "number"}),
    ])


def _presentation_spec() -> PresentationSpec:
    return PresentationSpec(
        title="Status Jaringan",
        sections=[
            SlideSection(heading="Overview", content=["Semua normal"], image_ref="topology.png"),
            SlideSection(heading="Traffic", table={"headers": ["Device", "Mbps"], "rows": [["R1", "12.5"]]}),
        ],
        conclusion="Tidak ada insiden minggu ini.",
    )


class TestDocxGenerator(unittest.TestCase):

    def test_generates_valid_docx(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "out.docx"
            counts = generate_docx(_doc_spec(), output)

            self.assertTrue(output.exists())
            self.assertGreater(output.stat().st_size, 0)
            self.assertEqual(counts["headings"], 1)
            self.assertEqual(counts["tables"], 1)

            from docx import Document
            reopened = Document(str(output))
            self.assertGreaterEqual(len(reopened.paragraphs), 1)
            self.assertEqual(len(reopened.tables), 1)
            self.assertEqual(len(reopened.tables[0].rows), 3)  # header + 2 rows


class TestPdfGenerator(unittest.TestCase):

    def test_generates_valid_pdf(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "out.pdf"
            counts = generate_pdf(_doc_spec(), output)

            self.assertTrue(output.exists())
            data = output.read_bytes()
            self.assertTrue(data.startswith(b"%PDF-"))
            self.assertEqual(counts["tables"], 1)


class TestXlsxGenerator(unittest.TestCase):

    def test_generates_valid_xlsx(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "out.xlsx"
            counts = generate_xlsx(_sheet_spec(), output)

            self.assertTrue(output.exists())
            self.assertEqual(counts["worksheets"], 1)
            self.assertEqual(counts["rows"], 2)

            from openpyxl import load_workbook
            workbook = load_workbook(output)
            sheet = workbook["Traffic"]
            self.assertEqual(sheet["A1"].value, "Device")
            self.assertEqual(sheet["B1"].value, "Mbps")
            self.assertEqual(sheet["A2"].value, "R1")


class TestPptxGenerator(unittest.TestCase):

    def test_generates_valid_pptx(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "out.pptx"
            counts = generate_pptx(_presentation_spec(), output)

            self.assertTrue(output.exists())
            self.assertEqual(counts["sections"], 2)
            self.assertEqual(counts["tables"], 1)
            self.assertEqual(counts["image_placeholders"], 1)

            from pptx import Presentation
            reopened = Presentation(str(output))
            # title slide + 2 sections + 1 conclusion slide
            self.assertEqual(len(reopened.slides), 4)


class TestTextGenerators(unittest.TestCase):

    def test_generate_markdown(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "out.md"
            generate_markdown(_doc_spec(), output)
            text = output.read_text(encoding="utf-8")
            self.assertIn("# Laporan Jaringan", text)
            self.assertIn("| Device | Status |", text)

    def test_generate_txt(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "out.txt"
            generate_txt(_doc_spec(), output)
            text = output.read_text(encoding="utf-8")
            self.assertIn("Laporan Jaringan", text)


class TestCsvGenerator(unittest.TestCase):

    def test_generate_csv(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "out.csv"
            spec = CsvSpec(headers=["A", "B"], rows=[["1", "2"], ["3", "4"]])
            counts = generate_csv(spec, output)

            self.assertEqual(counts["rows"], 2)
            with open(output, newline="", encoding="utf-8") as handle:
                rows = list(csv.reader(handle))
            self.assertEqual(rows[0], ["A", "B"])
            self.assertEqual(rows[1], ["1", "2"])


if __name__ == "__main__":
    unittest.main()