"""core/file_processing/processor.py - Processor Registry sederhana:
image -> OCR + vision | pdf -> text layer, scan -> OCR | office/text -> ekstraksi |
archive -> ekstrak aman, rekursif. Semua hasil dari pemrosesan nyata; tanpa truncate."""
from __future__ import annotations

import logging
import tempfile
from pathlib import Path

from core.file_processing import archive as arc
from core.file_processing.models import (
    E_EXTRACTION, E_OCR, E_RESOURCE, E_SECURITY, E_UNSUPPORTED, E_VISION,
    FILE_COMPLETED, FILE_FAILED, FILE_PARTIAL, FILE_UNSUPPORTED, ProcessingResult,
)

logger = logging.getLogger("aira.file_processing")

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif"}
OFFICE_EXT = {".docx", ".xlsx", ".pptx"}
TEXT_EXT = {".txt", ".md", ".csv", ".json", ".yaml", ".yml", ".log", ".rsc", ".conf", ".cfg", ".ini",
            ".xml", ".html", ".py", ".sh", ".tsv", ".rtf"}
MIME = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp",
        ".bmp": "image/bmp", ".tif": "image/tiff", ".tiff": "image/tiff", ".gif": "image/gif"}


class FileProcessor:
    def __init__(self, ocr=None, vision=None, limits=None):
        self._ocr, self._vision = ocr, vision
        self.limits = limits or arc.ArchiveLimits()

    @property
    def ocr(self):
        if self._ocr is None:
            from core.file_processing.ocr import TesseractOCR
            self._ocr = TesseractOCR()
        return self._ocr

    @property
    def vision(self):
        if self._vision is None:
            from core.file_processing.vision_llm import LLMVisionProvider
            self._vision = LLMVisionProvider()
        return self._vision

    def process(self, path, display_name: str = None) -> ProcessingResult:
        path = Path(path)
        name = display_name or path.name
        result = ProcessingResult(source=name, type=self._kind(path, name))
        try:
            self._file(result, path, name, 0)
        except Exception as exc:  # tidak ditelan: dicatat sebagai error nyata
            logger.exception("processing gagal: %s", name)
            result.error(E_EXTRACTION, f"{type(exc).__name__}: {exc}", name)
        return result.finalize()

    # ---------------------------------------------------------------- dispatch
    def _kind(self, path: Path, name: str) -> str:
        ext = Path(name).suffix.lower()
        if arc.archive_kind(name):
            return "archive"
        if ext in IMAGE_EXT:
            return "image"
        if ext == ".pdf":
            return "pdf"
        if ext in OFFICE_EXT:
            return "document"
        if ext in TEXT_EXT or self._looks_text(path):
            return "text"
        return "unknown"

    @staticmethod
    def _looks_text(path: Path) -> bool:
        try:
            head = path.read_bytes()[:8192] if path.stat().st_size else b""
            head.decode("utf-8")
            return b"\x00" not in head
        except (UnicodeDecodeError, OSError):
            return False

    def _file(self, result, path: Path, rel: str, depth: int):
        kind = self._kind(path, rel)
        entry = {"path": rel, "type": kind, "status": FILE_COMPLETED, "processors": [], "detail": ""}
        result.files.append(entry)
        before = len(result.errors)
        try:
            if kind == "archive":
                self._archive(result, path, rel, depth, entry)
            elif kind == "image":
                self._image(result, path, rel, entry)
            elif kind == "pdf":
                self._pdf(result, path, rel, entry)
            elif kind == "document":
                self._office(result, path, rel, entry)
            elif kind == "text":
                result.text.append({"source_file": rel, "page": None, "kind": "text",
                                    "text": path.read_text(encoding="utf-8", errors="replace")})
                entry["processors"].append("text")
            else:
                entry["status"] = FILE_UNSUPPORTED
                result.error(E_UNSUPPORTED, "Format file tidak didukung.", rel)
        except MemoryError:
            entry["status"] = FILE_FAILED
            result.error(E_RESOURCE, "Memori tidak cukup untuk memproses file ini.", rel)
        if kind != "archive" and entry["status"] == FILE_COMPLETED and len(result.errors) > before:
            entry["status"] = FILE_PARTIAL if entry["processors"] else FILE_FAILED

    # ---------------------------------------------------------------- archive
    def _archive(self, result, path, rel, depth, entry):
        if depth >= self.limits.max_depth:
            entry["status"] = FILE_FAILED
            result.error(E_SECURITY, f"Archive bersarang melewati kedalaman {self.limits.max_depth}.", rel)
            return
        with tempfile.TemporaryDirectory(prefix="aira_extract_") as tmp:  # cleanup otomatis
            report = arc.extract_archive(path, Path(tmp), self.limits)
            for code, msg, p in report.errors:
                result.error(code, msg, f"{rel}!{p}" if p else rel)
            entry["processors"].append("archive")
            entry["detail"] = f"{len(report.entries)} file diekstrak"
            result.metadata.setdefault("archives", []).append({"path": rel, "extracted": len(report.entries)})
            if report.fatal and not report.entries:
                entry["status"] = FILE_FAILED
                return
            if report.errors or report.fatal:
                entry["status"] = FILE_PARTIAL
            for e in report.entries:
                self._file(result, e["abs"], f"{rel}!/{e['path']}", depth + 1)

    # ---------------------------------------------------------------- image
    def _image(self, result, path, rel, entry):
        from PIL import Image, ImageSequence
        ok = False
        try:
            img = Image.open(path)
            frames = [f.copy() for f in ImageSequence.Iterator(img)] if getattr(img, "n_frames", 1) > 1 else [img]
            width, height = img.size
        except Image.DecompressionBombError as exc:
            entry["status"] = FILE_FAILED
            result.error(E_RESOURCE, f"Gambar terlalu besar untuk diproses aman: {exc}", rel)
            return
        except Exception as exc:
            entry["status"] = FILE_FAILED
            result.error(E_EXTRACTION, f"Gambar tidak bisa dibuka: {exc}", rel)
            return

        for n, frame in enumerate(frames, 1):
            try:
                o = self.ocr.ocr_image(frame)
                result.ocr.append({"source_file": rel, "page": n, **o})
                ok = True
            except Exception as exc:
                result.error(E_OCR, f"OCR gagal ({type(exc).__name__}): {exc}", rel)
        if ok:
            entry["processors"].append("ocr")

        vision = self._run_vision(result, path, rel, width, height)
        if vision == "available":
            entry["processors"].append("vision")
        entry["detail"] = f"{width}x{height}"
        if not entry["processors"]:
            entry["status"] = FILE_FAILED

    def _run_vision(self, result, path, rel, width, height) -> str:
        from core.vision.models import ImageMetadata, VisionResult
        meta = ImageMetadata(name=Path(rel).name, mime_type=MIME.get(Path(rel).suffix.lower(), "image/*"),
                             size=path.stat().st_size, width=width, height=height)
        try:
            v = self.vision.analyze({"kind": "workspace_path", "absolute_path": str(path)}, meta)
        except Exception as exc:
            v = VisionResult(status="error", reason=f"{type(exc).__name__}: {exc}")
        result.images.append({"source_file": rel, "width": width, "height": height, "vision": v.to_dict()})
        if v.status == "error":
            result.error(E_VISION, f"Vision gagal: {v.reason}", rel)
        return v.status

    # ---------------------------------------------------------------- pdf
    def _pdf(self, result, path, rel, entry):
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            total = len(reader.pages)
        except ImportError:
            entry["status"] = FILE_FAILED
            return result.error(E_EXTRACTION, "pypdf belum terpasang.", rel)
        except Exception as exc:
            entry["status"] = FILE_FAILED
            return result.error(E_EXTRACTION, f"PDF tidak bisa dibaca: {exc}", rel)

        pdfium = None
        for n, page in enumerate(reader.pages, 1):
            try:
                text = (page.extract_text() or "").strip()
            except Exception as exc:
                text = ""
                result.error(E_EXTRACTION, f"Halaman {n}: {exc}", rel)
            if text:
                result.text.append({"source_file": rel, "page": n, "kind": "pdf_text", "text": text})
                if "text_layer" not in entry["processors"]:
                    entry["processors"].append("text_layer")
                continue
            try:  # halaman scan -> OCR
                if pdfium is None:
                    import pypdfium2
                    pdfium = pypdfium2.PdfDocument(str(path))
                img = pdfium[n - 1].render(scale=2.5).to_pil()
                o = self.ocr.ocr_image(img)
                result.ocr.append({"source_file": rel, "page": n, **o})
                if "ocr" not in entry["processors"]:
                    entry["processors"].append("ocr")
            except ImportError:
                result.error(E_OCR, f"Halaman {n} berupa scan dan pypdfium2 belum terpasang untuk merender.", rel)
            except Exception as exc:
                result.error(E_OCR, f"OCR halaman {n} gagal ({type(exc).__name__}): {exc}", rel)
        entry["detail"] = f"{total} halaman"

    # ---------------------------------------------------------------- office
    def _office(self, result, path, rel, entry):
        ext = Path(rel).suffix.lower()
        try:
            if ext == ".docx":
                from docx import Document
                d = Document(str(path))
                parts = [p.text for p in d.paragraphs]
                for t in d.tables:
                    parts += [" | ".join(c.text for c in r.cells) for r in t.rows]
                items = [(None, "\n".join(parts))]
            elif ext == ".xlsx":
                from openpyxl import load_workbook
                wb = load_workbook(str(path), read_only=True, data_only=True)
                items = [(ws.title, "\n".join(" | ".join("" if c is None else str(c) for c in row)
                                              for row in ws.iter_rows(values_only=True))) for ws in wb.worksheets]
            else:
                from pptx import Presentation
                items = [(i, " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame))
                         for i, s in enumerate(Presentation(str(path)).slides, 1)]
        except ImportError as exc:
            entry["status"] = FILE_FAILED
            return result.error(E_EXTRACTION, f"Library {exc.name} belum terpasang.", rel)
        except Exception as exc:
            entry["status"] = FILE_FAILED
            return result.error(E_EXTRACTION, f"Dokumen tidak bisa dibaca: {exc}", rel)
        for page, text in items:
            result.text.append({"source_file": rel, "page": page, "kind": "document_text", "text": text})
        entry["processors"].append("document_text")
