"""agents/rei/artifact_tools.py — tool LLM untuk ArtifactEngine & Attachment."""
import json
from pathlib import Path
from typing import Any, Optional

from core.artifacts import (
    ArtifactType, DocumentSpec, SpreadsheetSpec, PresentationSpec, CsvSpec, get_artifact_engine,
)
from core.artifacts.models import DOCUMENT_TYPES, SPREADSHEET_TYPES, PRESENTATION_TYPES
from core.attachments import get_attachment_engine

_TEXT_EXT = {".txt", ".md", ".csv", ".json", ".yaml", ".yml", ".log"}


def create_artifact(artifact_type: str, spec: Any, name: Optional[str] = None,
                    session_id: Optional[str] = None, **_ignored) -> dict:
    if isinstance(spec, str):
        try:
            spec = json.loads(spec)
        except ValueError:
            return {"success": False, "tool": "create_artifact", "error": "spec harus JSON object."}
    if not isinstance(spec, dict):
        return {"success": False, "tool": "create_artifact", "error": "spec harus object."}

    try:
        kind = ArtifactType.coerce(artifact_type).value
    except ValueError as exc:
        return {"success": False, "tool": "create_artifact", "error": str(exc)}

    engine = get_artifact_engine()
    kw = dict(session_id=session_id, name=name, source="ai_agent")

    if kind in DOCUMENT_TYPES:
        result = engine.create_document(DocumentSpec.from_dict(spec), kind, **kw)
    elif kind in SPREADSHEET_TYPES:
        s = CsvSpec.from_dict(spec) if kind == "csv" and "worksheets" not in spec else SpreadsheetSpec.from_dict(spec)
        result = engine.create_spreadsheet(s, kind, **kw)
    elif kind in PRESENTATION_TYPES:
        result = engine.create_presentation(PresentationSpec.from_dict(spec), kind, **kw)
    else:
        return {"success": False, "tool": "create_artifact", "error": f"Tipe '{kind}' tidak didukung."}

    if not result.success:
        return {"success": False, "tool": "create_artifact", "error": "; ".join(result.errors)}

    a = result.artifact
    return {
        "success": True, "tool": "create_artifact",
        "artifact_id": a.id, "name": a.name, "type": a.artifact_type, "size": a.size,
        "download_url": f"/api/artifacts/{a.id}/download",
        "note": "Beri user link Markdown [nama](download_url) PERSIS seperti ini.",
    }


def _extract(path: Path, max_chars: int):
    ext = path.suffix.lower()
    try:
        if ext in _TEXT_EXT:
            text = path.read_text(encoding="utf-8", errors="replace")
        elif ext == ".pdf":
            from pypdf import PdfReader
            text = "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
        elif ext == ".docx":
            from docx import Document
            d = Document(str(path))
            text = "\n".join(p.text for p in d.paragraphs)
            for t in d.tables:
                for row in t.rows:
                    text += "\n" + " | ".join(c.text for c in row.cells)
        elif ext == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(str(path), read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets:
                lines.append(f"## {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    lines.append(" | ".join("" if c is None else str(c) for c in row))
                    if sum(len(x) for x in lines) > max_chars:
                        break
            text = "\n".join(lines)
        elif ext == ".pptx":
            from pptx import Presentation
            text = "\n".join(
                f"[Slide {i}] " + " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
                for i, s in enumerate(Presentation(str(path)).slides, 1)
            )
        else:
            return None, False, f"Format '{ext}' belum bisa diekstrak."
    except ImportError as exc:
        return None, False, f"Library '{exc.name}' belum terpasang (pip install)."
    except Exception as exc:
        return None, False, f"Gagal membaca file: {exc}"
    return text[:max_chars], len(text) > max_chars, None


def list_attachments(session_id: Optional[str] = None, **_ignored) -> dict:
    items = get_attachment_engine().list(session_id=session_id) if session_id else []
    return {"success": True, "tool": "list_attachments", "count": len(items),
            "attachments": [{"id": a.id, "name": a.name, "mime_type": a.mime_type, "size": a.size} for a in items]}


def read_attachment(attachment: str, max_chars: Any = 6000, session_id: Optional[str] = None, **_ignored) -> dict:
    engine = get_attachment_engine()
    items = engine.list(session_id=session_id) if session_id else []
    key = (attachment or "").strip().lower()
    match = next((a for a in items if a.id.lower() == key or a.name.lower() == key), None)
    if match is None:
        names = ", ".join(a.name for a in items) or "(belum ada)"
        return {"success": False, "tool": "read_attachment",
                "error": f"Attachment '{attachment}' tidak ditemukan. Yang ada: {names}."}

    if match.mime_type.startswith("image/"):
        img = (match.metadata or {}).get("image") or {}
        return {"success": True, "tool": "read_attachment", "name": match.name,
                "note": "Ini gambar; vision provider belum terpasang, isi gambar tidak bisa dibaca. "
                        f"Dimensi: {img.get('width')}x{img.get('height')}."}

    ref = engine.get_content_reference(match.id, session_id=session_id)
    if not ref or ref.get("kind") != "workspace_path":
        return {"success": False, "tool": "read_attachment", "error": "File tidak tersedia untuk dibaca."}

    try:
        limit = max(500, min(int(max_chars), 20000))
    except (TypeError, ValueError):
        limit = 6000

    text, truncated, error = _extract(Path(ref["absolute_path"]), limit)
    if error:
        return {"success": False, "tool": "read_attachment", "error": error}
    return {"success": True, "tool": "read_attachment", "name": match.name,
            "content": text, "truncated": truncated}


ARTIFACT_TOOLS = {"create_artifact": create_artifact, "list_attachments": list_attachments,
                  "read_attachment": read_attachment}
ARTIFACT_TOOL_CATEGORY = {"create_artifact": "filesystem", "list_attachments": "filesystem",
                          "read_attachment": "filesystem"}

ARTIFACT_TOOL_SCHEMAS = [
    {"type": "function", "function": {
        "name": "create_artifact",
        "description": (
            "Membuat FILE dokumen nyata yang bisa diunduh user: docx, pdf, xlsx, pptx, csv, markdown, txt. "
            "Pakai saat user minta 'buatkan laporan/dokumen/spreadsheet/presentasi'. "
            "Bentuk spec: docx/pdf/markdown/txt -> {\"title\":\"..\",\"blocks\":[{\"type\":\"heading\",\"text\":\"..\",\"level\":1},"
            "{\"type\":\"paragraph\",\"text\":\"..\"},{\"type\":\"list\",\"items\":[..]},{\"type\":\"table\",\"headers\":[..],\"rows\":[[..]]}]}; "
            "xlsx/csv -> {\"worksheets\":[{\"name\":\"Sheet1\",\"headers\":[..],\"rows\":[[..]]}]}; "
            "pptx -> {\"title\":\"..\",\"sections\":[{\"heading\":\"..\",\"content\":[\"poin\"]}],\"conclusion\":\"..\"}. "
            "Jumlah kolom tiap row HARUS sama dengan headers."),
        "parameters": {"type": "object", "properties": {
            "artifact_type": {"type": "string", "description": "docx|pdf|xlsx|pptx|csv|markdown|txt"},
            "spec": {"type": "object", "description": "Struktur konten sesuai tipe."},
            "name": {"type": "string", "description": "Judul/nama file (opsional)."}},
            "required": ["artifact_type", "spec"]}}},
    {"type": "function", "function": {
        "name": "list_attachments",
        "description": "Melihat file yang di-upload user ke chat ini.",
        "parameters": {"type": "object", "properties": {}, "required": []}}},
    {"type": "function", "function": {
        "name": "read_attachment",
        "description": "Membaca ISI file yang di-upload user (txt/md/csv/json/pdf/docx/xlsx/pptx). WAJIB dipakai kalau user bertanya soal file yang dilampirkan - jangan mengarang isinya.",
        "parameters": {"type": "object", "properties": {
            "attachment": {"type": "string", "description": "id atau nama file."},
            "max_chars": {"type": "integer", "default": 6000}},
            "required": ["attachment"]}}},
]